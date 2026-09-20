"""Convert FL DOE 'In-Field and Out-of-Field' teaching xls/xlsx files (Alachua only)
into data/processed/infield_outfield_alachua.json.

Shape: { "<school_year>": { "<school_number>": {
    "school_name": ..., "total_classes": ..., "in_field": ..., "out_of_field": ...,
    "pct_in_field": ..., "pct_out_of_field": ...
} } }

Note: excludes district-level 'DISTRICT TOTALS' row (school # 0) -- keep only real schools.
"""
import json, re, glob, os
import openpyxl
import xlrd

def year_label(fname):
    m = re.search(r"IFOFFTeach(\d{2})(\d{2})", os.path.basename(fname))
    return f"20{m.group(1)}-{m.group(2)}"

def convert_xlsx(fname):
    wb = openpyxl.load_workbook(fname, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = next(i for i, r in enumerate(rows) if r and str(r[0]).strip() == "District #")
    data_rows = rows[hdr_idx + 1:]
    return parse_rows(data_rows)

def convert_xls(fname):
    wb = xlrd.open_workbook(fname)
    ws = wb.sheet_by_index(0)
    all_rows = [ws.row_values(i) for i in range(ws.nrows)]
    hdr_idx = next(i for i, r in enumerate(all_rows) if r and str(r[0]).strip() == "District #")
    data_rows = all_rows[hdr_idx + 1:]
    return parse_rows(data_rows)

def parse_rows(data_rows):
    schools = {}
    for r in data_rows:
        if not r or len(r) < 9:
            continue
        district_name = r[1]
        if district_name != "ALACHUA":
            continue
        school_num_raw = r[2]
        school_name = r[3]
        if not school_num_raw or str(school_num_raw).strip() in ("0", "0.0"):
            continue  # skip DISTRICT TOTALS row
        school_num = str(int(float(school_num_raw))).zfill(4)

        total, in_field, out_field, pct_in, pct_out = r[4], r[5], r[6], r[7], r[8]
        schools[school_num] = {
            "school_name": school_name,
            "total_classes": total,
            "in_field": in_field,
            "out_of_field": out_field,
            "pct_in_field": pct_in,
            "pct_out_of_field": pct_out,
        }
    return schools

def main():
    raw_dir = os.path.join(os.path.dirname(__file__), "raw", "infield")
    out_dir = os.path.join(os.path.dirname(__file__), "processed")
    os.makedirs(out_dir, exist_ok=True)

    files = sorted(glob.glob(os.path.join(raw_dir, "IFOFFTeach*.xls*")))
    result = {}
    for f in files:
        yr = year_label(f)
        print(f"{yr}: {os.path.basename(f)}")
        if f.endswith(".xlsx"):
            result[yr] = convert_xlsx(f)
        else:
            result[yr] = convert_xls(f)
        print(f"  -> {len(result[yr])} Alachua schools")

    out_path = os.path.join(out_dir, "infield_outfield_alachua.json")
    with open(out_path, "w") as fh:
        json.dump(result, fh, indent=2)
    print(f"\nwrote {out_path}")

if __name__ == "__main__":
    main()
